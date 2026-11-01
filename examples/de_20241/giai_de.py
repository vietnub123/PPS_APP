r"""Giải đề 20241 bằng dữ liệu, guide và bộ tính sẵn trong repo.

Không tải tài liệu, không dùng bộ giải tham chiếu bên ngoài. Các phép kiểm
bổ sung chỉ dùng số hữu tỷ, phép nhân/cộng đa thức và quét dấu trên bảng.
Chạy từ thư mục gốc: .venv\Scripts\python.exe examples\de_20241\giai_de.py
"""
from pathlib import Path
from fractions import Fraction as F
from decimal import Decimal, localcontext, ROUND_FLOOR, ROUND_CEILING
from hashlib import sha256
from contextlib import redirect_stdout
import csv
import io
import json
import math
import sys

ROOT=Path(__file__).resolve().parents[2]
sys.path.insert(0,str(ROOT))
from openpyxl import load_workbook
from pps.common import Report, data_pairs, set_display_digits
from pps.gui import dispatch
from pps.worked_report import ReportBuilder
from pps.math_view import export_html
from pps.inverse_model import difference_table, evaluate, derivative, polynomial_range
from exam_cli import main as cli_main

OUT=Path(__file__).resolve().parent


def mul(a,b):
    c=[F(0)]*(len(a)+len(b)-1)
    for i,x in enumerate(a):
        for j,y in enumerate(b):c[i+j]+=x*y
    return c


def plus(a,b):
    c=[F(0)]*max(len(a),len(b))
    for i,v in enumerate(a):c[i]+=v
    for i,v in enumerate(b):c[i]+=v
    return c


def shifted(a,offset):
    c=[a[-1]]
    for v in reversed(a[:-1]):c=plus(mul(c,[offset,F(1)]),[v])
    return c


def dec(v):
    if isinstance(v,F):return Decimal(v.numerator)/Decimal(v.denominator)
    return Decimal(str(v))


def bound(v,upper=False):
    """Làm tròn ra ngoài khi in một đầu mút của bao khoảng."""
    with localcontext() as c:
        c.prec=100
        return format(dec(v).quantize(Decimal('1e-12'),rounding=ROUND_CEILING if upper else ROUND_FLOOR),'f')


def interval_text(pair):return '['+bound(pair[0])+', '+bound(pair[1],True)+']'


def save_task(name,tool,method,values,x='',y='',*,html=True):
    task={'tool':tool,'method':method,'values':{'digits':'12','display_digits':'12',**values},'x':x,'y':y}
    path=OUT/(name+'.json');path.write_text(json.dumps(task,ensure_ascii=False,indent=2),encoding='utf-8')
    r=dispatch(tool,method,task['values'],x,y)
    path.with_suffix('.txt').write_text(r.text(),encoding='utf-8')
    if html:export_html(r,path.with_suffix('.html'))
    return task,r


def markdown(report,path):
    lines=['# '+report.title,'']
    def table(index):
        title,headers,rows=report.tables[index]
        cell=lambda v:str(v).replace('|','\\|').replace('\n','<br>')
        lines.extend(['**'+title+'**','','| '+' | '.join(map(cell,headers))+' |',
                      '| '+' | '.join('---' for _ in headers)+' |'])
        lines.extend('| '+' | '.join(map(cell,row))+' |' for row in rows)
        lines.append('')
    def visit(nodes,level=2):
        for node in nodes:
            kind=node['kind']
            if kind=='columns':visit(node['children'],level)
            elif kind in ('section','panel'):
                if node.get('title'):lines.extend(['#'*level+' '+node['title'],''])
                visit(node['children'],min(level+1,6))
            elif kind=='details':
                lines.extend(['<details>','<summary>'+node['title']+'</summary>',''])
                visit(node['children'],min(level+1,6));lines.extend(['</details>',''])
            elif kind=='text':lines.extend([node['text'],''])
            elif kind=='formula':lines.extend(['$$',node['latex'],'$$',''])
            elif kind=='table':table(node['index'])
    visit(report.presentation['worked_layout'])
    lines.extend(['Bản trình bày công thức ngoại tuyến: [HTML](Bai_giai_de_20241.html).',
                  'Quá trình thao tác và lệnh chạy lại: [nhật ký](examples/de_20241/Nhat_ky_thuc_hien.md).',''])
    path.write_text('\n'.join(lines),encoding='utf-8')


def build_outputs(*,html=True):
    set_display_digits(12)
    OUT.mkdir(exist_ok=True)
    workbook=ROOT/'kienthuc'/'20241GKtest2.xlsx'
    book=load_workbook(workbook,read_only=True,data_only=True)
    active=book.active.title;sheet_rows={s.title:list(s.values) for s in book};book.close()
    rows=sheet_rows['Sheet3'];xy=data_pairs([r[:2] for r in rows])
    x=[a for a,b in xy];y=[b for a,b in xy]
    fx=list(map(lambda v:F(str(v)),x));fy=list(map(lambda v:F(str(v)),y))
    xbar,ybar=F(str(rows[0][3])),F(str(rows[0][5]))
    h=fx[1]-fx[0]
    assert active=='Sheet3' and (xbar,ybar)==(F(4),F(6))
    assert len(x)==106 and all(b-a==h for a,b in zip(fx,fx[1:])) and h==F('0.117')
    assert [(r[1],r[2]) for r in sheet_rows['Sheet1'][1:]]==xy
    with (OUT/'du_lieu_Sheet3.csv').open('w',encoding='utf-8-sig',newline='') as f:
        writer=csv.writer(f);writer.writerow(['i','dong_Excel','x_i','y_i'])
        writer.writerows((i,i+2,a,b) for i,(a,b) in enumerate(xy))
    tasks=[]
    a_task,a_report=save_task('A_Horner','Đạo hàm & Horner','Horner',
        {'coef':'3.5 -2.78 5.63 0 -15.28 1.54','target':'2.57','order':'3'},html=html)
    tasks.append(('A_Horner',0))
    coefs=list(map(F,['3.5','-2.78','5.63','0','-15.28','1.54']));c=F('2.57')
    work=coefs.copy();horner=[]
    for r in range(4):
        acc=F(0);trace=[]
        for v in work:acc=acc*c+v;trace.append(acc)
        horner.append({'r':r,'input':work,'accumulated':trace,'remainder':acc,'value':math.factorial(r)*acc})
        assert abs(float(a_report.tables[r][2][-1][1])-float(acc))<1e-9
        work=trace[:-1]
    assert horner[2]['value']==F('1054.696646') and horner[3]['value']==F('1249.3386')
    c_task,c_report=save_task('C_Stirling_9','Nội suy đa thức','Stirling',
        {'target':'4','order':'1','count':'9'},x,y,html=html);tasks.append(('C_Stirling_9',0))
    ids=sorted(sorted(range(len(x)),key=lambda i:(abs(fx[i]-xbar),i))[:9]);assert ids==list(range(22,31))
    xx=[fx[i] for i in ids];yy=[fy[i] for i in ids];origin=xx[4];tau=(xbar-origin)/h
    columns,_=difference_table(xx,yy);stirling=[];q=[F(0)]*9
    for m in range(9):
        if m==0:sigma=yy[4];basis=[F(1)];source='y_0'
        elif m%2:
            r=(m-1)//2;sigma=(columns[m][4-r-1]+columns[m][4-r])/2;basis=[F(0),F(1)]
            for j in range(1,r+1):basis=mul(basis,[-F(j*j),F(0),F(1)])
            source=f'(Δ^{m} y_{-r-1} + Δ^{m} y_{-r})/2'
        else:
            r=m//2;sigma=columns[m][4-r];basis=[F(0),F(0),F(1)]
            for j in range(1,r):basis=mul(basis,[-F(j*j),F(0),F(1)])
            source=f'Δ^{m} y_{-r}'
        beta=sigma/math.factorial(m);q=plus(q,[beta*v for v in basis])
        stirling.append({'m':m,'source':source,'sigma':sigma,'beta':beta,'basis':basis,
                         'value':beta*evaluate(basis,tau),'derivative':beta*evaluate(derivative(basis),tau)/h})
    value=sum((a['value'] for a in stirling),F(0));slope=sum((a['derivative'] for a in stirling),F(0))
    assert abs(float(value)-float(c_report.value(4)))<1e-11
    assert abs(float(slope)-float(c_report.value.deriv()(4)))<1e-11
    brackets=[i for i in range(len(x)-1) if (fy[i]-ybar)*(fy[i+1]-ybar)<0]
    assert brackets==[14,23,62,82]
    choices=[(list(range(11,20)),'Tiến'),(ids,'Tiến'),(list(range(58,67)),'Tiến'),(list(range(78,87)),'Lùi')]
    branches=[]
    for k,((selected,direction),b) in enumerate(zip(choices,brackets),1):
        task,r=save_task(f'D{k}_Lap_{"tien" if direction=="Tiến" else "lui"}',
            'Nội suy ngược','Lặp có cận co',{'target':'6','a':str(x[b]),'b':str(x[b+1]),
                'selected_indices':selected,'direction':direction,'tol':'1e-6','residual_tol':'1e-6','max_iter':'100'},x,y,html=html)
        d=r.presentation['inverse_solution'];assert d['failure'] is None,d['failure']
        cert=d['iteration']['certificate'];assert cert['verified'] and cert['q']<1
        assert x[b]<=r.value<=x[b+1]
        assert d['iteration']['history'][-1]['position_ok'] and d['iteration']['history'][-1]['residual_ok']
        branches.append(d);tasks.append((f'D{k}_Lap_{"tien" if direction=="Tiến" else "lui"}',0))
    # C và D2 là đúng một đa thức, khác gốc: t_c=t_d−4.
    assert shifted(q,F(-4))==branches[1]['model']['polynomial']
    fail_task,fail_report=save_task('D4_Thu_tien_chua_du_chung_nhan','Nội suy ngược','Lặp có cận co',
        {**tasks_to_values(branches[3]),'direction':'Tiến'},x,y,html=html)
    assert fail_report.presentation['inverse_solution']['status']=='certification_unavailable'
    tasks.append(('D4_Thu_tien_chua_du_chung_nhan',1))

    report=Report('Bài giải giữa kỳ MI3042 — học kỳ 20241, Đề 2, Sheet3')
    w=ReportBuilder(report,display_digits=12);num=w.number;tex=w.tex
    with w.group('section','Dữ kiện và cách đọc đề'):
        w.text('Nguồn duy nhất: ảnh kienthuc/20241_MI3042_GK.jpg, workbook kienthuc/20241GKtest2.xlsx, guide và bộ tính trong repo. Không dùng Internet, lời giải hay dữ liệu từ bên ngoài.')
        w.text('Chọn Sheet3: ghi chú viết tay trên ảnh ghi “Đề 2 sheet 3”, workbook cũng đang mở Sheet3. Sheet1 lặp lại dữ liệu ở B:C nhưng không có x̄,ȳ; Sheet2 là bộ dữ liệu khác, x̄=4,ȳ=5. Sheet3 có x̄=4,ȳ=6 tại D1,F1.')
        w.table('Dữ liệu sử dụng',['số cặp','x đầu','x cuối','h','x̄','ȳ'],[[106,fx[0],fx[-1],h,xbar,ybar]])
        w.text('Đề gồm: (a) thuật toán Horner và đạo hàm cấp 2,3; (b) thuật toán Stirling 2N+1 mốc; (c) trích 9 mốc, tính giá trị và đạo hàm tại 4; (d) thuật toán lặp tìm x tại mức 6. Câu d không cho bậc/dung sai: bài này chọn rõ 9 mốc cho mỗi nhánh và ε_x=ε_y=10⁻⁶ theo guide, tối đa 100 vòng. Không coi các lựa chọn ấy là dữ kiện được in trong đề.')

    with w.group('section','Câu a. Horner và đạo hàm cấp 2, cấp 3 tại c=2.57'):
        w.formula(r'P(x)=3.5x^5-2.78x^4+5.63x^3+0x^2-15.28x+1.54','P(x)=3.5x^5−2.78x^4+5.63x^3+0x²−15.28x+1.54.')
        w.text('Nhập hệ số theo bậc giảm: 3.5, −2.78, 5.63, 0, −15.28, 1.54. Hệ số x² bằng 0 phải được giữ.')
        w.text('Thuật toán: đặt A là dãy hệ số giảm dần. Với r=0,…,min(k,N): b₀=A₀; lần lượt b_j=A_j+c·b_(j−1). Lấy s_r là phần tử cuối; trả P^(r)(c)=r!·s_r; thay A bằng dãy b bỏ phần tử cuối rồi làm lượt tiếp theo. Nếu r>N thì đạo hàm bằng 0.')
        w.formula(r'P(x)=s_0+(x-c)s_1+(x-c)^2s_2+\cdots,\qquad P^{(r)}(c)=r!s_r',
                  'Chia Horner liên tiếp cho x−c tạo các hệ số Taylor s_r; vì thế phải nhân r! để lấy đạo hàm.')
        for start in (0,2):
            with w.group('columns'):
                for a in horner[start:start+2]:
                    r=a['r']
                    with w.group('panel',f'Lượt {r+1}: tính s_{r}'):
                        w.table('Từng phép nhân–cộng Horner',['j','A_j','b_j=A_j+2.57·b_(j−1)'],
                                ((j,a0,b0) for j,(a0,b0) in enumerate(zip(a['input'],a['accumulated']))))
                        w.formula(rf's_{{{r}}}={tex(a["remainder"])},\quad P^{{({r})}}(2.57)={r}!s_{{{r}}}={tex(a["value"])}',
                                  f's_{r}={num(a["remainder"])}; P^({r})(2.57)={r}!s_{r}={num(a["value"])}.')
        w.formula(r'P^{\prime\prime}(2.57)=1054.696646,\qquad P^{\prime\prime\prime}(2.57)=1249.3386',
                  'Đáp số câu a: P″(2.57)=1054.696646; P‴(2.57)=1249.3386.')

    with w.group('section','Câu b. Thuật toán Stirling với 2N+1 mốc'):
        w.text('1. Sắp x tăng và giữ cặp (x,y); kiểm tra x phân biệt, cách đều h. Cần ít nhất 2N+1 mốc. Tìm mốc gần x̄; lấy cửa sổ liên tiếp có 2N+1 mốc quanh đó, điều chỉnh cửa sổ nếu chạm đầu bảng. Gọi mốc giữa là x₀, đánh số j=−N,…,N; kiểm tra vị trí t̄=(x̄−x₀)/h.')
        w.text('2. Dựng bảng: Δ⁰y_j=y_j; với r=1,…,2N, tính Δʳy_j=Δ^(r−1)y_(j+1)−Δ^(r−1)y_j. Giữ đúng chỉ số mốc đầu mỗi ô.')
        w.text('3. Cộng các hạng lẻ từ trung bình hai sai phân đối xứng, các hạng chẵn từ sai phân giữa; giai thừa và tích cơ sở được ghi trong công thức dưới. Tích rỗng bằng 1.')
        w.formula(r'Q_{2N}(t)=y_0+S_{\rm le}(t)+S_{\rm chan}(t),\qquad P(x)=Q_{2N}\!\left(\frac{x-x_0}{h}\right)',
                  'Q_2N(t)=y₀+S_lẻ(t)+S_chẵn(t); P(x)=Q_2N((x−x₀)/h).')
        w.formula(r'S_{\rm le}(t)=\sum_{r=0}^{N-1}\frac{\Delta^{2r+1}y_{-r-1}+\Delta^{2r+1}y_{-r}}{2(2r+1)!}\,t\prod_{j=1}^{r}(t^2-j^2)',
                  'Hạng lẻ 2r+1: trung bình Δ^(2r+1)y_(-r−1), Δ^(2r+1)y_(-r); chia (2r+1)!; nhân t·tích_(j=1)^r(t²−j²).')
        w.formula(r'S_{\rm chan}(t)=\sum_{r=1}^{N}\frac{\Delta^{2r}y_{-r}}{(2r)!}\,t^2\prod_{j=1}^{r-1}(t^2-j^2)',
                  'Hạng chẵn 2r: Δ^(2r)y_(-r)/(2r)! nhân t²·tích_(j=1)^(r−1)(t²−j²).')
        w.text('4. Thế t̄ để tính giá trị. Nếu hỏi đạo hàm cấp một, lấy Q′(t̄)/h. Kiểm tra đa thức đi qua các mốc được chọn; kiểm tra này không phải cận sai số với hàm thật.')

    with w.group('section','Câu c. Chọn 9 mốc và tính tại x̄=4'):
        w.text('Chọn N=4, bậc không quá 8. Mốc gần 4 nhất là x₀=4.042, cách 0.042; mốc 3.925 cách 0.075. Lấy bốn mốc mỗi phía: các dòng dữ liệu i=22,…,30, tương ứng dòng Excel 24,…,32. Không dùng toàn bộ 106 mốc cho câu yêu cầu chín mốc.')
        w.table('Chín mốc và bảng sai phân tiến',['j','x_j','y_j']+[f'Δ^{r}y_j' for r in range(1,9)],
                ([i-4,xx[i],yy[i]]+[columns[r][i] if i<len(columns[r]) else None for r in range(1,9)] for i in range(9)))
        w.formula(r't=\frac{x-4.042}{0.117},\qquad \bar t=-\frac{14}{39}\approx-0.358974358974',
                  't=(x−4.042)/0.117; tại x=4: t̄=−14/39≈−0.358974358974.')
        w.table('Sai phân dùng trong công thức',['m','nguồn σ_m','σ_m','β_m=σ_m/m!'],
                ((a['m'],a['source'],a['sigma'],a['beta']) for a in stirling),exact=True,text_columns=(1,))
        w.formula(r'Q_8(t)=5.3284-0.3247t-0.0173t^2+\frac{t(t^2-1)}{1200}+\frac{t^2(t^2-1)}{40000}',
                  'Q₈(t)=5.3284−0.3247t−0.0173t²+t(t²−1)/1200+t²(t²−1)/40000')
        w.formula(r'\qquad-\frac{t^2(t^2-1)(t^2-4)}{1200000}-\frac{t(t^2-1)(t^2-4)(t^2-9)}{25200000}',
                  '−t²(t²−1)(t²−4)/1200000−t(t²−1)(t²−4)(t²−9)/25200000')
        w.formula(r'\qquad+\frac{t^2(t^2-1)(t^2-4)(t^2-9)}{16800000}',
                  '+t²(t²−1)(t²−4)(t²−9)/16800000. Hạng cấp 5 bằng 0 vì hai sai phân cấp 5 đối xứng triệt tiêu.')
        w.text('Đặt B_m(t) là tích cơ sở tương ứng ở câu b, T_m=β_m B_m(t̄). Tính cả đóng góp đạo hàm theo x để kiểm tra hệ số 1/h:')
        w.table('Thế t̄=−14/39 vào từng hạng',['m','β_m','B_m(t̄)','T_m','β_m B′_m(t̄)/h'],
                ((a['m'],a['beta'],evaluate(a['basis'],tau),a['value'],a['derivative']) for a in stirling))
        w.formula(rf'P_8(4)=\sum_{{m=0}}^8 T_m\approx {tex(value)},\qquad P_8^\prime(4)=\frac{{Q_8^\prime(\bar t)}}{{0.117}}\approx {tex(slope)}',
                  f'P₈(4)≈{num(value)}; P₈′(4)≈{num(slope)}.')
        w.text('Kết quả câu c: f(4)≈5.442986424022; f′(4)≈−2.673281525950 theo mô hình Stirling chín mốc. Đề không cho cận đạo hàm hàm thật/sai số dữ liệu, nên không khẳng định tất cả chữ số in ra đều chính xác với f.')

    with w.group('section','Câu d. Lặp giải f(x)=ȳ=6'):
        w.text('Trước tiên quét dấu y_i−6 trên nguyên bảng. Có bốn cặp liên tiếp đổi dấu; dữ liệu không tạo một nhánh đơn điệu duy nhất. Mỗi khoảng được xử lý riêng. Nếu f liên tục thì mỗi khoảng có ít nhất một nghiệm; quét dấu không loại được nghiệm tiếp xúc ở nơi khác.')
        w.table('Các khoảng đổi dấu',['nhánh','i trái','a','f(a)−6','b','f(b)−6'],
                ((k+1,i,fx[i],fy[i]-ybar,fx[i+1],fy[i+1]-ybar) for k,i in enumerate(brackets)))
        w.text('Chọn chín mốc cho mỗi nhánh như bảng dưới để giữ bậc tối đa 8. Nhánh 2 dùng chính chín mốc của câu c, do đó Newton và Stirling mô tả đúng một đa thức trên nhánh ấy. Các tập khác được công bố rõ; không ghép thành một đa thức duy nhất đi qua cả 106 mốc.')
        w.table('Lựa chọn mốc, hướng và gốc',['nhánh','i dùng (từ 0)','miền mốc','hướng','gốc','T=J'],
                ((k+1,f'{d["selection"]["indices"][0]}…{d["selection"]["indices"][-1]}',
                  f'[{x[d["selection"]["indices"][0]]}, {x[d["selection"]["indices"][-1]]}]',
                  choices[k][1],d['model']['anchor'],str(tuple(map(str,d['iteration']['T'])))) for k,d in enumerate(branches)),text_columns=(1,2,3,5))
        w.text('Thuật toán: (1) dựng sai phân và chọn gốc đầu nếu tiến, gốc cuối nếu lùi; (2) lấy c_r=Δʳy_đầu hoặc ∇ʳy_cuối; (3) tách p(t)=c₀+d·t+H(t), d=c₁≠0; (4) t₀=(6−c₀)/d; (5) kiểm tra t₀∈J, Φ(J)⊆J và q<1; (6) lặp, đổi về x và kiểm tra đồng thời cận vị trí/phần dư; không đạt sau 100 vòng thì báo chưa đạt.')
        w.formula(r'p(t)=c_0+d\,t+H(t),\quad\Phi(t)=\frac{6-c_0-H(t)}{d},\quad t_{k+1}=\Phi(t_k),\quad x_k=x_{\rm goc}+0.117t_k',
                  'p(t)=c₀+d·t+H(t); Φ(t)=(6−c₀−H(t))/d; t_(k+1)=Φ(t_k); x_k=gốc+0.117t_k.')
        w.formula(r'H_{\rm tien}(t)=\sum_{r=2}^{8}\frac{c_r}{r!}\prod_{j=0}^{r-1}(t-j),\qquad H_{\rm lui}(t)=\sum_{r=2}^{8}\frac{c_r}{r!}\prod_{j=0}^{r-1}(t+j)',
                  'Tiến: H=sum_(r=2)^8 c_r/r!*t(t−1)…(t−r+1); lùi: dùng t(t+1)…(t+r−1).')
        w.table('Hệ số sai phân; c₁ chính là d',['r']+[f'nhánh {k+1}' for k in range(4)],
                ([r]+[d['model']['coefficients'][r] for d in branches] for r in range(9)))
        failed_seed=fail_report.presentation['inverse_solution']['iteration']['t0']
        w.text(f'Nhánh 4 chọn lùi: nếu dùng tiến với cùng chín mốc thì t₀={num(failed_seed)} ngoài J=[4,5], chưa áp dụng được chứng nhận với khởi tạo mặc định. Đổi hướng là lựa chọn được công bố trước lượt tính chính; với lùi, t₀∈[−4,−3]. Đây không phải thiếu mốc, và không kẹp t₀ vào biên.')
        w.text('Chặn hội tụ dùng công cụ hữu tỷ trong pps/inverse_model.py: chặn từng tích/hạng của H để bao Φ(J); chặn đa thức Φ′ trên cả J bằng đạo hàm/đơn điệu hoặc đổi gốc và tổng trị tuyệt đối hệ số. Các đầu mút thập phân dưới đây được làm tròn ra ngoài 12 chữ số; q được giữ dạng phân số đúng, không lấy max từ vài điểm mẫu.')
        w.table('Kiểm tra co trên toàn đoạn',['nhánh','t₀','bao Φ(J) ⊆ J','bao Φ′(J)','q<1'],
                ((k+1,d['iteration']['t0'],interval_text(d['iteration']['certificate']['image']),
                  interval_text(d['iteration']['certificate']['derivative_image']),str(d['iteration']['certificate']['q'])) for k,d in enumerate(branches)),text_columns=(2,3,4))
        w.formula(r'B_k=0.117\frac{q}{1-q}|t_k-t_{k-1}|\leq10^{-6},\qquad r_k=|p(t_k)-6|\leq10^{-6}',
                  'Dừng khi cả B_k=0.117*q/(1−q)*|t_k−t_(k−1)|≤10⁻⁶ và r_k=|p(t_k)−6|≤10⁻⁶.')
        for k,d in enumerate(branches,1):
            it=d['iteration'];model=d['model'];hist=it['history'];cert=it['certificate']
            with w.group('panel',f'Nhánh {k}: [{x[brackets[k-1]]}, {x[brackets[k-1]+1]}] — {choices[k-1][1].lower()}'):
                w.text(f'Gốc={num(model["anchor"])}; h=0.117; c₀={num(model["coefficients"][0])}; d={num(model["coefficients"][1])}; t₀={num(it["t0"])}; q={cert["q"]}≈{num(cert["q"])}.')
                data0=[0,it['t0'],None,model['anchor']+h*it['t0'],abs(evaluate(model['polynomial'],it['t0'])-ybar),None,'Khởi tạo']
                looprows=[data0]+[[a['iteration'],a['t_new'],a['high_sum'],a['x_new'],a['residual_y'],a['error_bound_x'],
                                  'Đạt' if a['status']=='accepted' else 'Tiếp tục'] for a in hist]
                w.table('Toàn bộ vòng lặp; H dùng giá trị t ở dòng trước',['k','t_k','H(t_(k−1))','x_k','r_k','B_k','dừng'],looprows,text_columns=(6,))
                first,last=hist[0],hist[-1]
                w.text(f'Vòng 1: H(t₀)={num(first["high_sum"])}; t₁=({num(ybar-model["coefficients"][0])}−({num(first["high_sum"])}))/({num(model["coefficients"][1])})={num(first["t_new"])}; x₁={num(first["x_new"])}.')
                w.text(f'Vòng cuối {len(hist)}: B={num(last["error_bound_x"])}≤10⁻⁶ và r={num(last["residual_y"])}≤10⁻⁶, nên dừng. Nghiệm mô hình x≈{num(d["result"]["x_hat"])}.','success')
                with w.details('Từng số hạng cấp 2…8 trong tất cả vòng'):
                    w.text('T_r=c_r/r! nhân tích tiến/lùi tại t_(k−1). Mỗi hàng cộng T₂+…+T₈ để được H; không bỏ số hạng khác 0 hoặc làm tròn trước khi lặp.')
                    w.table('Đóng góp của từng cấp',['k']+[f'T_{r}' for r in range(2,9)]+['H'],
                            ([a['iteration']]+[term['value'] for term in a['term_contributions']]+[a['high_sum']] for a in hist))
                    w.table('Chặn từng hạng H trên J',['r','cận dưới','cận trên'],
                            ((a['order'],a['lower'],a['upper']) for a in cert['term_ranges']),exact=True)
        w.table('Kết quả câu d theo bốn mô hình cục bộ',['nhánh','x gần đúng','vòng','cận B cuối','phần dư cuối'],
                ((k+1,d['result']['x_hat'],d['iterations_done'],d['iteration']['history'][-1]['error_bound_x'],
                  d['iteration']['history'][-1]['residual_y']) for k,d in enumerate(branches)))
        w.text('Các giá trị khoảng 2.728733; 3.760763; 8.352579; 10.641567 là nghiệm gần đúng tìm được. Nhánh gần x̄=4 là nhánh 2; cùng đa thức với câu c nhưng đang tìm x tại mức 6, không phải tính tại x=4.')

    with w.group('section','Kết luận và phạm vi sai số'):
        w.text('Câu a: P″(2.57)=1054.696646; P‴(2.57)=1249.3386. Câu b: thuật toán Stirling đã nêu đủ chọn mốc, sai phân, chỉ số, tích cơ sở và đổi đạo hàm. Câu c: f(4)≈5.442986424022, f′(4)≈−2.673281525950 theo P₈. Câu d: bốn nhánh đã có bảng vòng lặp và kiểm tra dừng riêng.')
        w.text('Cận B là cận giải phương trình của từng đa thức nội suy. Đề chỉ cho bảng y bốn chữ số thập phân, không cho cận đạo hàm/sai số dữ liệu của f; vì vậy chưa có cận tổng sai số với nghiệm hàm thật, chưa chứng nhận toàn bộ sai số máy bằng số học khoảng. Các chữ số thừa trong bảng để truy phép tính, không phải lời hứa chính xác tương ứng.')
        w.text('Toàn bộ dữ liệu, JSON từng câu, bản xuất của bộ tính và mã tái lập đặt tại examples/de_20241. Quá trình đọc đề, phân biệt sheet, các lựa chọn và kiểm tra được ghi riêng trong Nhat_ky_thuc_hien.md.')
    w.finish()
    if html:export_html(report,ROOT/'Bai_giai_de_20241.html')
    markdown(report,ROOT/'Bai_giai_de_20241.md')
    (OUT/'Bai_giai_van_ban.txt').write_text(report.text(),encoding='utf-8')
    logs=[]
    for name,expected in tasks:
        with redirect_stdout(io.StringIO()):code=cli_main([str(OUT/(name+'.json'))])
        assert code==expected,(name,code)
        logs.append(f'{name}.json: exit={code}, expected={expected}, OK')
    checks=['Đúng Sheet3 đang chọn và ghi chú ảnh; D1=4,F1=6.',
            '106 cặp hữu hạn; h=0.117 bằng số hữu tỷ tại mọi bước; Sheet1 B:C lặp lại Sheet3 A:B.',
            'Horner hữu tỷ trùng các lượt Horner của app; đạo hàm cấp2/cấp3 đúng giá trị thập phân ở câu a.',
            'Stirling hữu tỷ cộng theo công thức đối xứng trùng giá trị/đạo hàm của bộ tính repo (<1e-11).',
            'Đa thức Newton nhánh2 bằng chính xác Q_Stirling(t−4), kiểm bằng toàn bộ hệ số hữu tỷ.',
            'Bốn nhánh đều có chứng nhận q<1 và đạt đồng thời hai ngưỡng; không dùng bộ giải ngoài repo.']
    sources=['Guide_di_thi_nhanh.md','Huong_dan_dung_khi_thi.md','kienthuc/20241_MI3042_GK.jpg',
             'kienthuc/20241GKtest2.xlsx','pps/interpolation.py','pps/inverse_solution.py','pps/inverse_model.py',
             'pps/worked_report.py','pps/inverse_export.py']
    meta={'sources':{s:sha256((ROOT/s).read_bytes()).hexdigest() for s in sources},'sheet':'Sheet3',
          'count':106,'xbar':4,'ybar':6,'h':'0.117','c_indices':ids,
          'horner_second':str(horner[2]['value']),'horner_third':str(horner[3]['value']),
          'stirling_value':str(value),'stirling_derivative':str(slope),
          'roots':[{'indices':d['selection']['indices'],'method':choices[k][1],
                    'value':d['result']['x_hat'],'q':str(d['iteration']['certificate']['q']),
                    'iterations':d['iterations_done'],'last_B':str(d['iteration']['history'][-1]['error_bound_x']),
                    'last_residual':str(d['iteration']['history'][-1]['residual_y'])} for k,d in enumerate(branches)],
          'checks':checks,'cli':logs}
    (OUT/'ket_qua_va_nguon.json').write_text(json.dumps(meta,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({k:meta[k] for k in ['count','xbar','ybar','horner_second','horner_third','roots']},ensure_ascii=False,indent=2))
    return report,meta


def tasks_to_values(d):
    return {'target':'6','a':str(float(d['interval'][0])),'b':str(float(d['interval'][1])),
            'selected_indices':d['selection']['indices'],'tol':'1e-6','residual_tol':'1e-6','max_iter':'100'}


if __name__=='__main__':build_outputs()
